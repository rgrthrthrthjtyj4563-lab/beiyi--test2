import docx
import pandas as pd
import os

prd_path = '/Users/lee/Documents/贝医/开发对账单系统/PRD/Smart Generation Tool for Pharma Compliance Statements _ PRD _ v1.0.docx'
config_path = '/Users/lee/Documents/贝医/开发对账单系统/Standard Configuration Table of Statement Billing Items.xlsx'

def read_docx(path):
    print(f"--- Reading DOCX: {os.path.basename(path)} ---")
    try:
        doc = docx.Document(path)
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        
        print("\n--- Tables in DOCX ---")
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                print(" | ".join(row_text))
            print("-" * 20)

    except Exception as e:
        print(f"Error reading docx: {e}")

def read_xlsx(path):
    print(f"\n--- Reading XLSX: {os.path.basename(path)} ---")
    try:
        # Read all sheets
        xls = pd.ExcelFile(path)
        for sheet_name in xls.sheet_names:
            print(f"\nSheet: {sheet_name}")
            df = pd.read_excel(xls, sheet_name=sheet_name)
            print(df.to_string())
    except Exception as e:
        print(f"Error reading xlsx: {e}")

if __name__ == "__main__":
    if os.path.exists(prd_path):
        read_docx(prd_path)
    else:
        print(f"File not found: {prd_path}")

    if os.path.exists(config_path):
        read_xlsx(config_path)
    else:
        print(f"File not found: {config_path}")

import importlib
import json
import os
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

warnings.filterwarnings("ignore", category=ResourceWarning, message=r"unclosed database.*")


def _clear_backend_modules() -> None:
    for name in ["backend.main", "backend.logic", "backend.manager", "backend.config_store"]:
        if name in sys.modules:
            del sys.modules[name]


def _load_fresh_app(temp_dir: str, legacy_payload: dict | None = None):
    db_path = Path(temp_dir) / "app.db"
    json_path = Path(temp_dir) / "statements_db.json"
    json_path.write_text(json.dumps(legacy_payload or {}, ensure_ascii=False), encoding="utf-8")

    os.environ["SOA_DB_PATH"] = str(db_path)
    os.environ["STATEMENTS_JSON_PATH"] = str(json_path)
    _clear_backend_modules()

    main = importlib.import_module("backend.main")
    logic = importlib.import_module("backend.logic")
    return main, logic


class StatementRegressionTests(unittest.TestCase):
    @staticmethod
    def _project_test_excel() -> Path:
        return Path(__file__).resolve().parents[2] / "PRD" / "测试业务使用数据.xlsx"

    def test_amount_conservation(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            record = main.create_statement(
                target_amount=100000,
                customer="守恒测试客户",
                period="2026-01",
                ratios="{}",
                rule_template_name="standard",
                template_name="default",
                mappings_json="{}",
                file=None,
            )

            generated = float(record.statement.summary.generated_amount)
            target = float(record.statement.target_amount)
            summed_items = sum(float(x.amount) for x in record.statement.items)
            self.assertAlmostEqual(generated, target, places=2)
            self.assertAlmostEqual(summed_items, target, places=2)

    def test_l4_ratio_boundary(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            main.put_allocation_rule(
                main.AllocationRuleTemplateUpdate(
                    template_name="standard",
                    strategy_name="equal_split_v1",
                    ratios={"L1": 0.6, "L2": 0.28, "L3": 0.10, "L4": 0.02},
                    l4_max_ratio=0.02,
                    ratio_warning_threshold=0.2,
                    fill_order=["L2", "L3", "L1"],
                )
            )
            record = main.create_statement(
                target_amount=50000,
                customer="边界测试客户",
                period="2026-02",
                ratios="{}",
                rule_template_name="standard",
                template_name="default",
                mappings_json="{}",
                file=None,
            )
            l4_ratio = float(record.statement.summary.layers["L4"]["ratio"])
            self.assertLessEqual(l4_ratio, 0.021)

    def test_parse_mapping(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            config_items = logic.load_config()
            self.assertGreater(len(config_items), 0)
            target_item_name = config_items[0].name

            df = pd.DataFrame(
                {
                    "年份": [2026, 2026, 2026],
                    "月份": [1, 1, 2],
                    "服务提供方": ["A", "B", "C"],
                    "外部字段": [10, 20, 30],
                }
            )
            excel_path = Path(td) / "biz.xlsx"
            df.to_excel(excel_path, index=False)

            parsed = logic.parse_business_data_preview(
                file_path=str(excel_path),
                period="2026-01",
                config_items=config_items,
                column_mapping={"外部字段": target_item_name},
            )

            self.assertEqual(parsed["rows_total"], 3)
            self.assertEqual(parsed["rows_used"], 2)
            self.assertIn(f"外部字段->{target_item_name}", parsed["matched_columns"])
            self.assertAlmostEqual(float(parsed["quantities"][target_item_name]), 30.0, places=6)

    def test_fixed_amount_over_target_blocks_with_explanation(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            config_items = logic.load_config()
            fixed_l1 = next(
                (i for i in config_items if i.level == "L1" and not i.can_simulate and i.price > 0),
                None,
            )
            self.assertIsNotNone(fixed_l1)

            df = pd.DataFrame(
                {
                    "年份": [2026],
                    "月份": [1],
                    "服务提供方": ["A"],
                    fixed_l1.name: [10000],
                }
            )
            excel_path = Path(td) / "fixed_over_target.xlsx"
            df.to_excel(excel_path, index=False)

            with self.assertRaises(ValueError) as ctx:
                logic.generate_smart_statement(
                    target_amount=1000,
                    customer="固定超额测试",
                    period="2026-01",
                    business_data_file=str(excel_path),
                )
            self.assertIn("固定费用超过目标金额", str(ctx.exception))

    def test_no_report_support_fallback_to_l3(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            stmt = logic.generate_smart_statement(
                target_amount=5000,
                customer="无报告回落测试",
                period="2026-01",
                business_data_file=None,
            )
            self.assertTrue(
                any("未检测到报告类业务支撑" in w for w in stmt.summary.warnings),
                msg=str(stmt.summary.warnings),
            )
            non_l1_levels = {item.level for item in stmt.items if item.level != "L1" and float(item.amount) > 0}
            self.assertTrue(non_l1_levels.issubset({"L3"}), msg=str(non_l1_levels))

    def test_non_simulated_item_uses_real_quantity(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            config_items = logic.load_config()
            l3_non_sim = next(
                (i for i in config_items if i.level == "L3" and not i.can_simulate and i.price > 0),
                None,
            )
            self.assertIsNotNone(l3_non_sim)

            df = pd.DataFrame(
                {
                    "年份": [2026],
                    "月份": [1],
                    "服务提供方": ["A"],
                    l3_non_sim.name: [2],
                }
            )
            excel_path = Path(td) / "l3_non_sim.xlsx"
            df.to_excel(excel_path, index=False)

            stmt = logic.generate_smart_statement(
                target_amount=5000,
                customer="不可模拟直算测试",
                period="2026-01",
                business_data_file=str(excel_path),
            )
            match = next((i for i in stmt.items if i.name == l3_non_sim.name and i.source == "业务真实数据"), None)
            self.assertIsNotNone(match)
            self.assertEqual(int(match.quantity), 2)
            self.assertAlmostEqual(float(match.amount), float(l3_non_sim.price) * 2, places=6)

    def test_tail_diff_is_balanced_into_l4_only(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            stmt = logic.generate_smart_statement(
                target_amount=100001,
                customer="尾差调平测试",
                period="2026-01",
                business_data_file=None,
            )
            self.assertAlmostEqual(float(stmt.summary.generated_amount), 100001.0, places=2)
            tail_rows = [i for i in stmt.items if i.level == "L4" and i.source == "尾差调平"]
            self.assertEqual(len(tail_rows), 1)
            self.assertAlmostEqual(float(tail_rows[0].amount), 1.0, places=6)

    def test_actual_selectable_respects_ratio_bounds(self):
        with tempfile.TemporaryDirectory() as td:
            main, logic = _load_fresh_app(td)
            items = main.config_store.list_billing_items(include_inactive=True)
            candidate = next((x for x in items if float(x.get("price", 0)) > 0), None)
            self.assertIsNotNone(candidate)

            main.config_store.update_billing_item(
                int(candidate["id"]),
                {
                    "quantity_mode": "ACTUAL_SELECTABLE",
                    "min_pick_ratio": 0.2,
                    "max_pick_ratio": 0.4,
                    "pick_priority": 1,
                    "status": "ACTIVE",
                },
            )

            df = pd.DataFrame(
                {
                    "年份": [2026],
                    "月份": [1],
                    "服务提供方": ["A"],
                    candidate["name"]: [100],
                }
            )
            excel_path = Path(td) / "selectable.xlsx"
            df.to_excel(excel_path, index=False)

            target_qty = 30
            target_amount = float(candidate["price"]) * target_qty
            stmt = logic.generate_smart_statement(
                target_amount=target_amount,
                customer="择取边界测试",
                period="2026-01",
                business_data_file=str(excel_path),
            )
            row = next((i for i in stmt.items if i.name == candidate["name"]), None)
            self.assertIsNotNone(row)
            self.assertEqual(row.quantity_mode, "ACTUAL_SELECTABLE")
            self.assertGreaterEqual(float(row.billed_qty), 20.0)
            self.assertLessEqual(float(row.billed_qty), 40.0)

    def test_runtime_mode_prefers_quantity_mode_over_billing_mode(self):
        with tempfile.TemporaryDirectory() as td:
            main, logic = _load_fresh_app(td)
            items = main.config_store.list_billing_items(include_inactive=True)
            candidate = next((x for x in items if x["level"] == "L3"), None)
            self.assertIsNotNone(candidate)

            updated = main.config_store.update_billing_item(
                int(candidate["id"]),
                {
                    "quantity_mode": "SIMULATED",
                    "billing_mode": "智能择取",
                    "can_simulate": 1,
                    "status": "ACTIVE",
                },
            )
            self.assertIsNotNone(updated)

            loaded = logic.load_config()
            runtime_item = next((x for x in loaded if int(x.id or 0) == int(candidate["id"])), None)
            self.assertIsNotNone(runtime_item)
            self.assertEqual(runtime_item.quantity_mode, "SIMULATED")

    def test_selectable_solver_never_overshoots_target(self):
        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            item = logic.BillingItem(
                name="可择取测试项",
                level="L1",
                unit="次",
                price=20.0,
                quantity_mode="ACTUAL_SELECTABLE",
                can_simulate=False,
                min_pick_ratio=0.0,
                max_pick_ratio=1.0,
            )
            entry = logic._SelectableEntry(item=item, actual_qty=1000)
            warnings = []
            amount, feasible = logic._select_actual_quantities([entry], 10011.0, warnings)
            self.assertTrue(feasible)
            self.assertLessEqual(float(amount), 10011.0)
            self.assertAlmostEqual(float(amount), 10000.0, places=6)

    def test_v2_snapshot_and_version_persisted(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            record = main.create_statement(
                target_amount=20000,
                customer="V2快照测试",
                period="2026-02",
                ratios="{}",
                rule_template_name="standard",
                template_name="default",
                mappings_json="{}",
                file=None,
            )
            self.assertEqual(record.engine_version, "v2")
            self.assertIsNotNone(record.rule_snapshot_version)
            self.assertIn("engine_version", record.statement.summary.calc_snapshot_json)
            self.assertGreaterEqual(len(record.statement.summary.item_decision_json), 1)

    def test_enabled_item_ids_gate_effective(self):
        with tempfile.TemporaryDirectory() as td:
            main, logic = _load_fresh_app(td)
            items = main.config_store.list_billing_items(include_inactive=True)
            enabled = [int(x["id"]) for x in items if x["level"] == "L3"][:2]
            l4_tail = next((int(x["id"]) for x in items if x["level"] == "L4"), None)
            if l4_tail is not None:
                enabled = enabled + [l4_tail]
            self.assertGreaterEqual(len(enabled), 1)

            main.put_allocation_rule(
                main.AllocationRuleTemplateUpdate(
                    template_name="standard",
                    strategy_name="equal_split_v1",
                    selection_strategy="priority_greedy_v1",
                    ratios={"L1": 0.0, "L2": 0.0, "L3": 1.0, "L4": 0.0},
                    l4_max_ratio=0.05,
                    ratio_warning_threshold=0.2,
                    fill_order=["L3", "L2", "L4"],
                    enabled_item_ids=enabled,
                    tail_diff_threshold=10000,
                    max_simulation_ratio=1.0,
                )
            )

            stmt = logic.generate_smart_statement(
                target_amount=10000,
                customer="启停门禁测试",
                period="2026-01",
                business_data_file=None,
                rule_template_name="standard",
            )
            names = {x["name"] for x in items if int(x["id"]) in set(enabled)}
            for row in stmt.items:
                if row.source == "尾差调平":
                    continue
                self.assertIn(row.name, names)

    def test_feasibility_endpoint_returns_reason_code(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            resp = main.precheck_statement(
                target_amount=10,
                customer="预检测试",
                period="2026-01",
                ratios="{}",
                rule_template_name="standard",
                template_name="default",
                mappings_json="{}",
                file=None,
            )
            self.assertIn("feasible", resp)
            self.assertIn("reason_code", resp)
            self.assertIn(resp["reason_code"], {"OK", "FIXED_EXCEEDS_TARGET", "SELECTABLE_INFEASIBLE", "TAIL_DIFF_EXCEEDED", "TAIL_ITEM_MISSING", "NO_ENABLED_ITEMS", "UNSPECIFIED"})

    def test_create_and_export_api_flow(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            record = main.create_statement(
                target_amount=32100,
                customer="集成测试客户",
                period="2026-03",
                ratios="{}",
                rule_template_name="standard",
                template_name="default",
                mappings_json="{}",
                file=None,
            )

            self.assertEqual(record.status, "DRAFT")
            export_resp = main.export_statement_file(record.id)
            self.assertEqual(export_resp.status_code, 200)
            self.assertTrue(Path(export_resp.path).exists())

            updated = main.get_statement(record.id)
            self.assertEqual(updated.status, "EXPORTED")
            self.assertTrue(any(h.get("action") == "EXPORT" for h in updated.history))

            try:
                Path(export_resp.path).unlink(missing_ok=True)
            except Exception:
                pass

    def test_export_template_clears_stale_rows(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            summary = main.StatementSummary(
                target_amount=33167.0,
                generated_amount=33167.0,
                diff=0.0,
                layers={
                    "L1": {"amount": 33160.0, "ratio": 0.9998},
                    "L2": {"amount": 0.0, "ratio": 0.0},
                    "L3": {"amount": 0.0, "ratio": 0.0},
                    "L4": {"amount": 7.0, "ratio": 0.0002},
                },
                warnings=[],
            )
            items = [
                main.StatementItem(level="L1", name="医院拜访", unit="元/次", price=20.0, quantity=1258.0, amount=25160.0, source="择取计费"),
                main.StatementItem(level="L1", name="工作组问卷调研分析报告", unit="元/人天", price=2000.0, quantity=4.0, amount=8000.0, source="择取计费"),
                main.StatementItem(level="L4", name="打印快递服务", unit="根据实际情况收取", price=1.0, quantity=7.0, amount=7.0, source="尾差调平"),
            ]
            statement = main.Statement(
                customer="测试6",
                period="2025-04",
                target_amount=33167.0,
                summary=summary,
                items=items,
            )
            record = main.manager.create(statement)
            export_resp = main.export_statement_file(record.id)
            path = Path(export_resp.path)
            self.assertTrue(path.exists())

            wb = load_workbook(path, data_only=False)
            ws = wb[wb.sheetnames[0]]
            total_rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "合计"]
            self.assertEqual(total_rows, [8], "导出模板中不应残留历史“合计”行")
            self.assertEqual(ws.cell(8, 7).value, "=SUM(G5:G7)")

            try:
                path.unlink(missing_ok=True)
            except Exception:
                pass

    def test_json_migration_once(self):
        with tempfile.TemporaryDirectory() as td:
            legacy_id = "legacy-001"
            legacy_payload = {
                legacy_id: {
                    "id": legacy_id,
                    "status": "DRAFT",
                    "created_at": "2026-02-11T00:00:00",
                    "updated_at": "2026-02-11T00:00:00",
                    "statement": {
                        "customer": "历史客户",
                        "period": "2026-01",
                        "target_amount": 100.0,
                        "summary": {
                            "target_amount": 100.0,
                            "generated_amount": 100.0,
                            "diff": 0.0,
                            "layers": {
                                "L1": {"amount": 100.0, "ratio": 1.0},
                                "L2": {"amount": 0.0, "ratio": 0.0},
                                "L3": {"amount": 0.0, "ratio": 0.0},
                                "L4": {"amount": 0.0, "ratio": 0.0},
                            },
                            "warnings": [],
                        },
                        "items": [
                            {
                                "level": "L1",
                                "name": "历史项",
                                "unit": "次",
                                "price": 100.0,
                                "quantity": 1.0,
                                "amount": 100.0,
                                "source": "系统固定",
                            }
                        ],
                    },
                    "history": [
                        {
                            "timestamp": "2026-02-11T00:00:00",
                            "action": "CREATED",
                            "note": "Initial generation",
                        }
                    ],
                }
            }

            main, _ = _load_fresh_app(td, legacy_payload=legacy_payload)
            migrated = main.get_statement(legacy_id)
            self.assertIsNotNone(migrated)
            self.assertEqual(migrated.statement.customer, "历史客户")

            with main.manager._connect() as conn:
                migration_rows = conn.execute(
                    "SELECT COUNT(1) AS c FROM statement_migrations WHERE migration_key = ?",
                    ("statements_json_to_sqlite_v1",),
                ).fetchone()
            self.assertEqual(int(migration_rows["c"]), 1)

    def test_list_billing_items_tolerates_dirty_rows(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            now = "2026-02-12T00:00:00"
            with main.config_store._connect() as conn:
                cursor = conn.execute(
                    """
                    INSERT INTO billing_items (
                        name, level, category, unit, price, billing_mode, pick_priority, require_business_data,
                        is_existing, can_simulate, must_use, is_enabled_default, quantity_mode, min_pick_ratio,
                        max_pick_ratio, report_required, tail_balance_eligible, allow_discount, status, sort_order,
                        effective_from, effective_to, billing_note, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        "",              # 脏值：空名称
                        "L9",            # 脏值：非法层级
                        "",
                        "",              # 脏值：空单位
                        "NaN",           # 脏值：非数值单价
                        "UNKNOWN",       # 脏值：非法计费模式
                        0,               # 脏值：异常优先级
                        1,               # 兼容字段（DB 非空）
                        0,
                        0,
                        0,
                        1,
                        "UNKNOWN",       # 脏值：非法 quantity_mode
                        2.0,             # 脏值：越界
                        -1.0,            # 脏值：越界 + 反序
                        0,
                        0,
                        0,
                        "ACTIVE",        # 兼容旧状态
                        0,
                        "",
                        "",
                        "",
                        now,
                        now,
                    ),
                )
                bad_id = int(cursor.lastrowid)
                conn.commit()

            items = main.list_billing_items(include_inactive=True)
            self.assertGreaterEqual(len(items), 1)
            bad = next((x for x in items if int(x.id) == bad_id), None)
            self.assertIsNotNone(bad)
            self.assertIn(bad.level, {"L1", "L2", "L3", "L4"})
            self.assertTrue(bool(bad.name.strip()))
            self.assertTrue(bool(bad.unit.strip()))
            self.assertIn(bad.status, {"启用", "停用"})
            self.assertGreaterEqual(float(bad.min_pick_ratio), 0.0)
            self.assertLessEqual(float(bad.max_pick_ratio), 1.0)

    def test_active_status_is_visible_when_include_inactive_false(self):
        with tempfile.TemporaryDirectory() as td:
            main, _ = _load_fresh_app(td)
            items = main.config_store.list_billing_items(include_inactive=True)
            candidate = next((x for x in items if str(x.get("status")) == "启用"), None)
            self.assertIsNotNone(candidate)
            main.config_store.update_billing_item(int(candidate["id"]), {"status": "ACTIVE"})

            active_only = main.config_store.list_billing_items(include_inactive=False)
            active_ids = {int(x["id"]) for x in active_only}
            self.assertIn(int(candidate["id"]), active_ids)

    def test_p0_4_strict_target_alignment_gate(self):
        test_file = self._project_test_excel()
        if not test_file.exists():
            self.skipTest(f"missing required fixture: {test_file}")

        with tempfile.TemporaryDirectory() as td:
            _, logic = _load_fresh_app(td)
            period = "2025-04"
            ranges = [(1000, 1100), (5000, 5100), (10000, 10100), (50000, 50100)]
            failed: list[tuple[int, str, str]] = []

            for start, end in ranges:
                for target in range(start, end + 1):
                    pre = logic.precheck_statement_feasibility(
                        target_amount=float(target),
                        customer="P0-4 Gate",
                        period=period,
                        business_data_file=str(test_file),
                        rule_template_name="standard",
                    )
                    if not pre["feasible"]:
                        failed.append((target, str(pre["reason_code"]), str(pre["message"])))
                        continue

                    stmt = logic.generate_smart_statement(
                        target_amount=float(target),
                        customer="P0-4 Gate",
                        period=period,
                        business_data_file=str(test_file),
                        rule_template_name="standard",
                    )
                    generated = float(stmt.summary.generated_amount)
                    if abs(generated - float(target)) > 1e-9:
                        failed.append((target, "DIFF_NOT_ZERO", f"generated={generated} target={target}"))

            self.assertEqual(
                len(failed),
                0,
                msg="; ".join([f"target={t},reason={r},msg={m}" for t, r, m in failed[:20]]),
            )


if __name__ == "__main__":
    unittest.main()

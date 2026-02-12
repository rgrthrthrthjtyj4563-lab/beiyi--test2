from fastapi import FastAPI, HTTPException, Body, UploadFile, File, Form, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from starlette.middleware.base import BaseHTTPMiddleware
from logic import (
    load_config,
    generate_smart_statement,
    precheck_statement_feasibility,
    precheck_statement_feasibility_lightweight,
    load_config_seed_payload,
    parse_business_data_preview,
    get_allocation_strategies,
    resolve_allocation_rule,
)
from template_parser import (
    generate_standard_input_template,
    parse_standard_template,
    parse_standard_template_l1,
    validate_standard_template,
)
from models import (
    Statement,
    BillingItem,
    BillingItemConfig,
    BillingItemConfigCreate,
    BillingItemConfigUpdate,
    BatchBillingItemUpdate,
    BatchBillingItemDelete,
    BusinessDataPreview,
    DataSourceRecord,
    FieldMappingTemplate,
    FieldMappingTemplateUpdate,
    AllocationRuleTemplate,
    AllocationRuleTemplateUpdate,
    AllocationRuleVersion,
    AllocationStrategyOption,
)
from manager import StatementManager, StatementRecord
from config_store import config_store
from openpyxl import load_workbook
import tempfile
import os
import shutil
import json
import anyio
from pathlib import Path
from datetime import datetime
from copy import copy
import logging
import inspect
import functools
from typing import Tuple, Dict, Optional

app = FastAPI(title="Smart Statement Generator")
manager = StatementManager()
logger = logging.getLogger(__name__)


class _ApiPrefixMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        path = request.scope.get("path") or ""
        prefixes = [p.strip() for p in os.getenv("SOA_API_PREFIXES", "/api").split(",") if p.strip()]
        for prefix in prefixes:
            if path == prefix or path.startswith(f"{prefix}/"):
                new_path = path[len(prefix):] or "/"
                request.scope["path"] = new_path
                request.scope["root_path"] = prefix
                request.scope["raw_path"] = new_path.encode()
                break
        return await call_next(request)

# Seed dynamic config table from existing excel config on first run.
config_store.seed_billing_items(load_config_seed_payload())
config_store.backfill_item_meta(load_config_seed_payload())

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.add_middleware(_ApiPrefixMiddleware)

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "PRD" / "Statement of Account.xlsx"
PREVIEW_TIMEOUT_SECONDS = 30
RUN_SYNC_SUPPORTS_CANCELLABLE = "cancellable" in inspect.signature(anyio.to_thread.run_sync).parameters


async def _run_parse_preview(**kwargs):
    file_path = kwargs.get('file_path', 'unknown')
    logger.info(f"开始解析文件: {file_path}, 超时设置: {PREVIEW_TIMEOUT_SECONDS}s")
    try:
        with anyio.fail_after(PREVIEW_TIMEOUT_SECONDS):
            parse_func = functools.partial(parse_business_data_preview, **kwargs)
            if RUN_SYNC_SUPPORTS_CANCELLABLE:
                result = await anyio.to_thread.run_sync(parse_func, cancellable=True)
            else:
                result = await anyio.to_thread.run_sync(parse_func)
        logger.info(f"文件解析完成: {file_path}, 结果行数: {result.get('rows_total', 0)}")
        return result
    except TimeoutError:
        logger.error(f"文件解析超时: {file_path}")
        raise HTTPException(status_code=408, detail=f"文件解析超时（>{PREVIEW_TIMEOUT_SECONDS}秒），请稍后重试或联系管理员")


def _load_preview_payload(preview_json: str):
    if not preview_json:
        return None
    try:
        payload = json.loads(preview_json)
    except Exception:
        return None
    if not isinstance(payload, dict):
        return None
    return payload


async def _resolve_business_data(
    preview_json: Optional[str],
    file: Optional[UploadFile],
    period: str,
    column_mapping: Dict,
) -> Tuple[Optional[Dict], Optional[str], Optional[Dict]]:
    """
    解析业务数据，优先使用前端预览数据，否则解析上传的文件
    返回: (quantities_dict, temp_file_path, preview_result)
    """
    # 1. 优先使用前端传入的预览数据
    preview_payload = _load_preview_payload(preview_json or "")
    if preview_payload:
        quantities = preview_payload.get("quantities")
        if isinstance(quantities, dict):
            logger.info(f"使用前端预览数据: {len(quantities)} 个计费项")
            return quantities, None, preview_payload

    # 2. 没有预览数据，解析上传的文件
    if file:
        suffix = os.path.splitext(file.filename or "")[1] if file.filename else ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            temp_file_path = tmp.name

        preview_result = await _run_parse_preview(
            file_path=temp_file_path,
            period=period,
            config_items=load_config(),
            column_mapping=column_mapping,
        )
        quantities = preview_result.get("quantities") if isinstance(preview_result, dict) else None

        if isinstance(quantities, dict):
            logger.info(f"文件解析完成: {len(quantities)} 个计费项")
            return quantities, temp_file_path, preview_result
        else:
            # 解析失败但文件已保存，返回文件路径以便后续使用
            return None, temp_file_path, preview_result

    return None, None, None


def _format_period_label(period: str) -> str:
    p = (period or "").strip().upper()
    if "-Q" in p:
        try:
            y, q = p.split("-Q")
            return f"{int(y)}年Q{int(q)}"
        except Exception:
            return period
    if "-" in p:
        try:
            y, m = p.split("-")
            return f"{int(y)}年{int(m)}月"
        except Exception:
            return period
    return period


def _level_label(level: str) -> str:
    mapping = {
        "L1": "L1系统收费层",
        "L2": "L2报告辅助加工",
        "L3": "L3增值服务",
        "L4": "L4其他服务",
    }
    return mapping.get(level, level)


def _unit_label(unit: str) -> str:
    if "/" in (unit or ""):
        return unit.split("/")[-1]
    return unit or ""


def _load_item_meta() -> dict:
    meta = {}
    db_items = config_store.list_billing_items(include_inactive=True)
    for item in db_items:
        name = str(item.get("name", "")).strip()
        if not name or name in meta:
            continue
        meta[name] = {
            "category": str(item.get("category", "") or ""),
            "reason": str(item.get("billing_note", "") or ""),
        }
    return meta


def _payload_dict(model) -> dict:
    if hasattr(model, "model_dump"):
        return model.model_dump()
    return model.dict()


def _fill_statement_template(statement: Statement) -> str:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb[wb.sheetnames[0]]

    ws["A2"] = statement.customer
    ws["B2"] = _format_period_label(statement.period)
    ws["C2"] = f"{statement.target_amount:,.2f}元"
    ws["D2"] = datetime.now().strftime("%Y-%m-%d")

    item_meta = _load_item_meta()
    items = statement.items

    start_row = 5
    base_detail_rows = 14
    if len(items) > base_detail_rows:
        extra_rows = len(items) - base_detail_rows
        ws.insert_rows(19, amount=extra_rows)
        # Keep inserted detail rows visually consistent with template.
        for offset in range(extra_rows):
            dst = 19 + offset
            src = 18
            for col in range(1, 10):
                ws.cell(dst, col)._style = copy(ws.cell(src, col)._style)

    detail_end_row = start_row + max(base_detail_rows, len(items)) - 1
    for row in range(start_row, detail_end_row + 1):
        for col in range(1, 10):
            ws.cell(row, col).value = None

    for idx, item in enumerate(items, start=start_row):
        meta = item_meta.get(item.name, {})
        ws.cell(idx, 1, _level_label(item.level))
        ws.cell(idx, 2, meta.get("category", ""))
        ws.cell(idx, 3, item.name)
        ws.cell(idx, 4, _unit_label(item.unit))
        ws.cell(idx, 5, item.quantity)
        ws.cell(idx, 6, item.price)
        ws.cell(idx, 7, item.amount)
        ws.cell(idx, 8, item.source)
        ws.cell(idx, 9, meta.get("reason", ""))

    total_row = start_row + len(items)
    clear_to = max(total_row + 4, 23)
    for row in range(total_row, clear_to + 1):
        for col in range(1, 10):
            ws.cell(row, col).value = None

    ws.cell(total_row, 1, "合计")
    if items:
        ws.cell(total_row, 7, f"=SUM(G{start_row}:G{total_row - 1})")
    else:
        ws.cell(total_row, 7, 0)

    layer_labels = {
        "L1": "L1系统收费层 小计",
        "L2": "L2报告辅助加工 小计",
        "L3": "L3增值服务 小计",
        "L4": "L4其他服务 小计",
    }
    for i, level in enumerate(["L1", "L2", "L3", "L4"], start=1):
        row = total_row + i
        layer = statement.summary.layers.get(level, {"amount": 0, "ratio": 0})
        ws.cell(row, 1, layer_labels[level])
        ws.cell(row, 7, layer.get("amount", 0))
        ws.cell(row, 9, f"占比 {float(layer.get('ratio', 0)) * 100:.2f}%")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        return tmp.name

@app.get("/config", response_model=list[BillingItem])
def get_config():
    return load_config()


@app.get("/settings/billing-items", response_model=list[BillingItemConfig])
def list_billing_items(include_inactive: bool = False):
    try:
        items = config_store.list_billing_items(include_inactive=include_inactive)
    except Exception as e:
        logger.exception("Failed to load billing items from store: %s", e)
        return []

    normalized: list[BillingItemConfig] = []
    for idx, item in enumerate(items):
        try:
            normalized.append(BillingItemConfig(**item))
        except Exception as e:
            logger.error("Skip invalid billing item at index %s: %s; raw=%s", idx, e, item)
            continue
    return normalized


@app.post("/settings/billing-items", response_model=BillingItemConfig)
def create_billing_item(payload: BillingItemConfigCreate):
    """创建计费项"""
    try:
        item = config_store.create_billing_item(_payload_dict(payload))
        if not item:
            raise HTTPException(status_code=400, detail="创建计费项失败")
        return BillingItemConfig(**item)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.put("/settings/billing-items/{item_id}", response_model=BillingItemConfig)
def update_billing_item(item_id: int, payload: BillingItemConfigUpdate):
    """更新计费项"""
    data = {k: v for k, v in _payload_dict(payload).items() if v is not None}
    item = config_store.update_billing_item(item_id, data)
    if not item:
        raise HTTPException(status_code=404, detail="计费项不存在")
    return BillingItemConfig(**item)


@app.post("/settings/billing-items/{item_id}/deactivate", response_model=BillingItemConfig)
def deactivate_billing_item(item_id: int):
    """停用计费项"""
    item = config_store.deactivate_billing_item(item_id)
    if not item:
        raise HTTPException(status_code=404, detail="计费项不存在")
    return BillingItemConfig(**item)


@app.post("/settings/billing-items/{item_id}/activate", response_model=BillingItemConfig)
def activate_billing_item(item_id: int):
    """启用计费项"""
    item = config_store.activate_billing_item(item_id)
    if not item:
        raise HTTPException(status_code=404, detail="计费项不存在")
    return BillingItemConfig(**item)


@app.delete("/settings/billing-items/{item_id}")
def delete_billing_item(item_id: int):
    """物理删除计费项"""
    success = config_store.delete_billing_item(item_id)
    if not success:
        raise HTTPException(status_code=404, detail="计费项不存在")
    return {"message": "计费项已删除", "id": item_id}


@app.post("/settings/billing-items/batch/update")
def batch_update_billing_items(payload: BatchBillingItemUpdate):
    """批量更新计费项"""
    try:
        updated_count = config_store.batch_update_billing_items(
            payload.item_ids, payload.updates
        )
        return {
            "message": f"成功更新 {updated_count} 个计费项",
            "updated_count": updated_count,
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/settings/billing-items/batch/delete")
def batch_delete_billing_items(payload: BatchBillingItemDelete):
    """批量物理删除计费项"""
    try:
        deleted_count = config_store.batch_delete_billing_items(payload.item_ids)
        return {
            "message": f"成功删除 {deleted_count} 个计费项",
            "deleted_count": deleted_count,
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

# Legacy endpoint - keeping it for compatibility or quick test
@app.post("/generate", response_model=Statement)
def generate_statement_legacy(
    target_amount: float = Body(..., embed=True),
    customer: str = Body(..., embed=True),
    period: str = Body(..., embed=True),
    ratios: dict = Body({'L1': 0.55, 'L2': 0.30, 'L3': 0.10, 'L4': 0.05}, embed=True)
):
    try:
        return generate_smart_statement(
            target_amount=target_amount,
            customer=customer,
            period=period,
            ratios=ratios,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/business-data/preview", response_model=BusinessDataPreview)
async def preview_business_data(
    period: str = Form(default=""),
    template_name: str = Form(default="default"),
    mappings_json: str = Form(default="{}"),
    file: UploadFile = File(...),
):
    temp_file_path = None
    file_size_mb = 0
    try:
        logger.info(f"收到文件预览请求: filename={file.filename}, period={period}")
        
        # 记录文件大小
        try:
            file.file.seek(0, 2)  # 移动到文件末尾
            file_size = file.file.tell()  # 获取位置（即文件大小）
            file.file.seek(0)  # 重置到开头
            file_size_mb = file_size / (1024 * 1024)
            logger.info(f"文件大小: {file_size_mb:.2f} MB")
            
            if file_size_mb > 100:
                logger.warning(f"文件过大: {file_size_mb:.2f} MB，可能导致解析缓慢")
        except Exception as size_err:
            logger.warning(f"无法获取文件大小: {size_err}")
        
        try:
            request_mapping = json.loads(mappings_json or "{}")
        except Exception:
            request_mapping = {}

        template_mapping = config_store.list_field_mappings(template_name=template_name)
        merged_mapping = {**template_mapping, **(request_mapping or {})}

        suffix = os.path.splitext(file.filename or "")[1] if file.filename else ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            temp_file_path = tmp.name

        parsed = await _run_parse_preview(
            file_path=temp_file_path,
            period=period,
            config_items=load_config(),
            column_mapping=merged_mapping,
        )
        
        logger.info(f"文件预览成功: rows_total={parsed.get('rows_total')}, rows_used={parsed.get('rows_used')}")
        return BusinessDataPreview(**parsed)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"文件预览失败: {e}")
        error_msg = str(e)
        if "timeout" in error_msg.lower() or "超时" in error_msg:
            raise HTTPException(
                status_code=408, 
                detail=f"文件解析超时（>{PREVIEW_TIMEOUT_SECONDS}秒）。文件大小: {file_size_mb:.1f}MB，建议稍后重试或联系管理员优化"
            )
        raise HTTPException(status_code=400, detail=f"文件解析失败: {error_msg}")
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)


@app.post("/business-data/preview-l1")
async def preview_business_data_l1(
    period: str = Form(default=""),
    preview_json: str = Form(default=""),
    file: UploadFile = File(default=None),
):
    temp_file_path = None
    try:
        preview_payload = _load_preview_payload(preview_json or "")
        if preview_payload and preview_payload.get("is_standard_template") and isinstance(preview_payload.get("quantities"), dict):
            config_items = load_config()
            l1_names = [item.name for item in config_items if item.require_business_data and item.level == "L1"]
            rows = []
            quantities = preview_payload.get("quantities") or {}
            for name in l1_names:
                rows.append({"name": name, "quantity": float(quantities.get(name, 0.0) or 0.0)})
            return {
                "success": True,
                "period": period or preview_payload.get("period", ""),
                "rows_total": preview_payload.get("rows_total", 0),
                "rows_used": preview_payload.get("rows_used", 0),
                "columns": ["名称", "数量"],
                "rows": rows,
                "is_standard_template": True,
                "level": "L1",
                "skipped_parse": True,
            }

        if not file:
            raise HTTPException(status_code=400, detail="缺少标准模板文件")

        suffix = os.path.splitext(file.filename or "")[1] if file.filename else ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            temp_file_path = tmp.name

        parsed = parse_standard_template_l1(
            file_path=temp_file_path,
            period=period,
            config_items=load_config(),
        )
        if not parsed.get("success"):
            raise HTTPException(status_code=400, detail=parsed.get("error", "解析失败"))
        return parsed
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"L1 模板解析失败: {e}")
        raise HTTPException(status_code=400, detail=f"L1 模板解析失败: {str(e)}")
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)


@app.post("/statements/feasibility")
async def precheck_statement(
    target_amount: float = Form(...),
    customer: str = Form(...),
    period: str = Form(...),
    ratios: str = Form(default="{}"),
    rule_template_name: str = Form(default="standard"),
    template_name: str = Form(default="default"),
    mappings_json: str = Form(default="{}"),
    preview_json: str = Form(default=""),
    file: UploadFile = File(default=None),
):
    temp_file_path = None
    try:
        logger.info(f"收到可行性检查请求: customer={customer}, period={period}, target={target_amount}")

        try:
            parsed_ratios = json.loads(ratios or "{}")
            ratios_dict = parsed_ratios if isinstance(parsed_ratios, dict) and parsed_ratios else None
        except Exception:
            ratios_dict = None

        try:
            request_mapping = json.loads(mappings_json or "{}")
        except Exception:
            request_mapping = {}
        template_mapping = config_store.list_field_mappings(template_name=template_name)
        merged_mapping = {**template_mapping, **(request_mapping or {})}

        # 使用公共函数解析业务数据
        preview_quantities, temp_file_path, preview_result = await _resolve_business_data(
            preview_json=preview_json,
            file=file,
            period=period,
            column_mapping=merged_mapping,
        )

        # Step 1: 轻量级检查 - 验证数据总量是否足够
        lightweight_result = precheck_statement_feasibility_lightweight(
            target_amount=target_amount,
            business_data_quantities=preview_quantities,
            config_items=load_config(),
        )

        if not lightweight_result.get("feasible"):
            # 轻量级检查不通过，直接返回错误
            logger.info(f"轻量级可行性检查不通过: {lightweight_result.get('reason_code')}")
            result = lightweight_result
        else:
            # Step 2: 轻量级检查通过，执行完整检查获取详细信息
            logger.info(f"轻量级检查通过，开始完整可行性检查")
            result = precheck_statement_feasibility(
                target_amount=target_amount,
                customer=customer,
                period=period,
                ratios=ratios_dict,
                business_data_file=temp_file_path,
                business_data_quantities=preview_quantities,
                column_mapping=merged_mapping,
                rule_template_name=rule_template_name,
            )
            # 保留轻量级检查的详细信息
            if "items" in lightweight_result:
                result["item_details"] = lightweight_result["items"]
                result["total_available"] = lightweight_result["total_available"]
                result["item_count"] = lightweight_result["item_count"]

        # 如果后端解析了文件，将解析结果返回给前端复用
        if preview_result and not preview_json:
            result["parse_result"] = {
                "quantities": preview_quantities,
                "rows_total": preview_result.get("rows_total", 0),
                "rows_used": preview_result.get("rows_used", 0),
                "matched_columns": preview_result.get("matched_columns", []),
                "unmatched_columns": preview_result.get("unmatched_columns", []),
            }

        logger.info(f"可行性检查完成: feasible={result.get('feasible')}, reason={result.get('reason_code')}")
        return result
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"可行性检查失败: {e}")
        raise HTTPException(status_code=500, detail=f"可行性检查失败: {str(e)}")
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)


# New Lifecycle Endpoints
@app.get("/statements", response_model=list[StatementRecord])
def list_statements():
    return manager.list()

@app.post("/statements", response_model=StatementRecord)
async def create_statement(
    target_amount: float = Form(...),
    customer: str = Form(...),
    period: str = Form(...),
    ratios: str = Form(default="{}"),
    rule_template_name: str = Form(default="standard"),
    template_name: str = Form(default="default"),
    mappings_json: str = Form(default="{}"),
    preview_json: str = Form(default=""),
    file: UploadFile = File(default=None),
):
    temp_file_path = None
    try:
        logger.info(f"收到创建对账单请求: customer={customer}, period={period}, target={target_amount}")

        if not file and not preview_json:
            raise HTTPException(status_code=400, detail="请上传业务明细文件")

        # Parse ratios
        try:
            parsed_ratios = json.loads(ratios or "{}")
            ratios_dict = parsed_ratios if isinstance(parsed_ratios, dict) and parsed_ratios else None
        except:
            ratios_dict = None

        try:
            request_mapping = json.loads(mappings_json or "{}")
        except Exception:
            request_mapping = {}
        template_mapping = config_store.list_field_mappings(template_name=template_name)
        merged_mapping = {**template_mapping, **(request_mapping or {})}

        file_name = file.filename if file else None

        # 使用公共函数解析业务数据
        preview_quantities, temp_file_path, preview_result = await _resolve_business_data(
            preview_json=preview_json,
            file=file,
            period=period,
            column_mapping=merged_mapping,
        )
        
        resolved_rule = resolve_allocation_rule(
            rule_template_name=rule_template_name,
            ratios=ratios_dict,
        )
        
        logger.info(f"开始生成对账单，策略: {resolved_rule.get('strategy_name')}")
        stmt = generate_smart_statement(
            target_amount, 
            customer, 
            period, 
            ratios_dict,
            business_data_file=temp_file_path,
            business_data_quantities=preview_quantities,
            column_mapping=merged_mapping,
            rule_template_name=rule_template_name,
        )
        
        record = manager.create(
            stmt,
            rule_template_name=str(resolved_rule.get("template_name", rule_template_name)),
            rule_template_version=int(resolved_rule.get("version", 1)),
            rule_strategy_name=str(resolved_rule.get("strategy_name", "")),
            engine_version=str((stmt.summary.calc_snapshot_json or {}).get("engine_version", "v2")),
            rule_snapshot_version=int((stmt.summary.calc_snapshot_json or {}).get("rule_snapshot_version", resolved_rule.get("version", 1))),
            rule_snapshot=resolved_rule,
        )
        
        source_payload = {
            "statement_id": record.id,
            "source_type": "BUSINESS_FILE",
            "source_name": file_name,
            "period": period,
            "rows_total": (preview_result or {}).get("rows_total", 0),
            "rows_used": (preview_result or {}).get("rows_used", 0),
            "matched_columns": len((preview_result or {}).get("matched_columns", [])),
            "unmatched_columns": len((preview_result or {}).get("unmatched_columns", [])),
        }
        config_store.create_data_source(source_payload)
        
        logger.info(f"对账单创建成功: id={record.id}, generated_amount={stmt.summary.generated_amount}")
        return record
                
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"生成对账单失败: {e}")
        import traceback
        error_detail = f"{str(e)}\n\n{traceback.format_exc()}"
        print(f"生成对账单失败:\n{error_detail}")
        raise HTTPException(status_code=500, detail=f"生成对账单失败: {str(e)}")
    finally:
        # Clean up temp file
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)

@app.get("/statements/{id}", response_model=StatementRecord)
def get_statement(id: str):
    record = manager.get(id)
    if not record:
        raise HTTPException(status_code=404, detail="Statement not found")
    return record


@app.get("/statements/{id}/sources", response_model=list[DataSourceRecord])
def get_statement_sources(id: str):
    record = manager.get(id)
    if not record:
        raise HTTPException(status_code=404, detail="Statement not found")
    rows = config_store.list_statement_sources(id)
    return [DataSourceRecord(**r) for r in rows]


@app.get("/settings/field-mappings", response_model=FieldMappingTemplate)
def get_field_mappings(template_name: str = "default"):
    mappings = config_store.list_field_mappings(template_name=template_name)
    return FieldMappingTemplate(template_name=template_name, mappings=mappings)


@app.put("/settings/field-mappings", response_model=FieldMappingTemplate)
def put_field_mappings(payload: FieldMappingTemplateUpdate):
    saved = config_store.replace_field_mappings(
        template_name=payload.template_name,
        mappings=payload.mappings,
    )
    return FieldMappingTemplate(template_name=payload.template_name, mappings=saved)


@app.get("/settings/allocation-rules", response_model=AllocationRuleTemplate)
def get_allocation_rule(template_name: str = "standard"):
    rule = config_store.get_allocation_rule(template_name=template_name)
    return AllocationRuleTemplate(**rule)


@app.put("/settings/allocation-rules", response_model=AllocationRuleTemplate)
def put_allocation_rule(payload: AllocationRuleTemplateUpdate):
    saved = config_store.upsert_allocation_rule(_payload_dict(payload))
    return AllocationRuleTemplate(**saved)


@app.get("/settings/allocation-rules/versions", response_model=list[AllocationRuleVersion])
def list_allocation_rule_versions(template_name: str = "standard", limit: int = 20):
    rows = config_store.list_allocation_rule_versions(template_name=template_name, limit=limit)
    return [AllocationRuleVersion(**r) for r in rows]


@app.get("/settings/allocation-strategies", response_model=list[AllocationStrategyOption])
def list_allocation_strategies():
    return [AllocationStrategyOption(**x) for x in get_allocation_strategies()]

@app.post("/statements/{id}/export")
def export_statement_file(id: str):
    record = manager.get(id)
    if not record:
        raise HTTPException(status_code=404, detail="Statement not found")

    statement = record.statement
    try:
        output_path = _fill_statement_template(statement)
        manager.update_status(id, "export", "Exported template attachment")
        return FileResponse(
            output_path,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename=f"Statement of Account_{statement.customer}_{statement.period}.xlsx"
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ========== 标准输入模板接口 ==========

@app.get("/templates/input/download")
def download_input_template():
    """下载标准化的业务数据输入模板

    模板特点：
    - 列名直接使用系统计费项名称，无需模糊匹配
    - 只需填写数量，无需其他字段
    - 支持多账期数据
    """
    temp_file_path = None
    try:
        # 在backend目录下创建临时文件，避免权限问题
        temp_dir = Path(__file__).parent / "data" / "temp"
        temp_dir.mkdir(parents=True, exist_ok=True)

        # 清理旧文件（保留1小时内的）
        _cleanup_old_temp_files(temp_dir, max_age_hours=1)

        temp_file_path = temp_dir / f"template_{os.urandom(4).hex()}.xlsx"
        generate_standard_input_template(str(temp_file_path))

        logger.info(f"模板生成成功: {temp_file_path}")

        # 使用 background 任务在文件发送后清理
        background_tasks = BackgroundTasks()
        background_tasks.add_task(_remove_file, path=str(temp_file_path))

        return FileResponse(
            path=str(temp_file_path),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename="business_data_template.xlsx",
            background=background_tasks
        )
    except Exception as e:
        # 出错时立即清理
        if temp_file_path and temp_file_path.exists():
            try:
                temp_file_path.unlink()
            except:
                pass
        logger.exception(f"生成模板失败: {e}")
        raise HTTPException(status_code=500, detail=f"生成模板失败: {str(e)}")


def _remove_file(path: str):
    """后台任务：删除临时文件"""
    try:
        file_path = Path(path)
        if file_path.exists():
            file_path.unlink()
            logger.info(f"临时文件已清理: {path}")
    except Exception as e:
        logger.warning(f"清理临时文件失败: {e}")


def _cleanup_old_temp_files(temp_dir: Path, max_age_hours: int = 1):
    """清理指定目录下超过指定时间的临时文件"""
    try:
        from datetime import datetime, timedelta
        cutoff_time = datetime.now() - timedelta(hours=max_age_hours)

        for file_path in temp_dir.glob("template_*.xlsx"):
            try:
                # 获取文件修改时间
                mtime = datetime.fromtimestamp(file_path.stat().st_mtime)
                if mtime < cutoff_time:
                    file_path.unlink()
                    logger.info(f"清理旧临时文件: {file_path}")
            except Exception as e:
                logger.warning(f"清理文件失败 {file_path}: {e}")
    except Exception as e:
        logger.warning(f"清理旧临时文件失败: {e}")


@app.post("/templates/input/validate")
async def validate_input_template(file: UploadFile = File(...)):
    """验证上传的文件是否符合标准模板格式"""
    temp_file_path = None
    try:
        suffix = os.path.splitext(file.filename or "")[1] if file.filename else ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            temp_file_path = tmp.name

        result = validate_standard_template(temp_file_path)
        return result
    except Exception as e:
        logger.exception(f"验证模板失败: {e}")
        raise HTTPException(status_code=400, detail=f"验证失败: {str(e)}")
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)


@app.post("/business-data/preview-standard", response_model=BusinessDataPreview)
async def preview_business_data_standard(
    period: str = Form(default=""),
    file: UploadFile = File(...),
):
    """使用标准模板解析业务数据（简化版，无需模糊匹配）

    相比 /business-data/preview 的优点：
    1. 列名直接匹配，速度快（通常 < 2秒）
    2. 无需字段映射配置
    3. 解析逻辑简单，不易出错
    4. 支持超大文件（因为无需逐行匹配）
    """
    temp_file_path = None
    try:
        logger.info(f"收到标准模板预览请求: filename={file.filename}, period={period}")

        # 记录文件大小
        try:
            file.file.seek(0, 2)
            file_size = file.file.tell()
            file.file.seek(0)
            file_size_mb = file_size / (1024 * 1024)
            logger.info(f"文件大小: {file_size_mb:.2f} MB")
        except Exception as size_err:
            logger.warning(f"无法获取文件大小: {size_err}")
            file_size_mb = 0

        suffix = os.path.splitext(file.filename or "")[1] if file.filename else ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            temp_file_path = tmp.name

        # 使用标准模板解析（快速模式）
        parsed = parse_standard_template(
            file_path=temp_file_path,
            period=period,
            config_items=load_config(),
        )

        if not parsed.get("success"):
            error_msg = parsed.get("error", "未知错误")
            logger.warning(f"标准模板解析失败: {error_msg}")
            raise HTTPException(status_code=400, detail=f"解析失败: {error_msg}")

        logger.info(f"标准模板解析成功: rows_total={parsed.get('rows_total')}, "
                   f"matched={len(parsed.get('matched_columns', []))}, "
                   f"quantities={len(parsed.get('quantities', {}))}")

        return BusinessDataPreview(**parsed)

    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"标准模板预览失败: {e}")
        raise HTTPException(status_code=400, detail=f"预览失败: {str(e)}")
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
